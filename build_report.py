#!/usr/bin/env python3
"""
Build Callaway Golf × Postie CRM Optimization reporting package.
Generates 5 PPTX one-sheeters + 1 HTML dashboard.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

OUTPUT_DIR = "/Users/taradonavanik/Postie API/Callaway/dashboard"
PPTX_DIR = os.path.join(OUTPUT_DIR, "one-sheeters")
os.makedirs(PPTX_DIR, exist_ok=True)

# ── Brand Colors ──────────────────────────────────────────────────────────────
NAVY      = RGBColor(0x16, 0x3D, 0x5D)
TEAL      = RGBColor(0x00, 0xC5, 0xC1)
ICE_BLUE  = RGBColor(0xEA, 0xF8, 0xFF)
LIGHT_BLUE= RGBColor(0xCA, 0xDC, 0xFC)
STEEL     = RGBColor(0x5A, 0x7A, 0x8A)
AMBER     = RGBColor(0xF0, 0xB8, 0x40)
AMBER_TINT= RGBColor(0xFF, 0xF3, 0xCD)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
DARK_TEXT = RGBColor(0x1E, 0x47, 0x68)
TEAL_LIGHT= RGBColor(0xCA, 0xDC, 0xFC)   # 5-10x iROAS row
GRAY_MID  = RGBColor(0x9B, 0xB0, 0xBB)

# ── Slide dimensions (widescreen 13.33" × 7.5") ──────────────────────────────
W = Inches(13.33)
H = Inches(7.5)
LEFT_W = Inches(13.33 * 0.35)   # 35%
RIGHT_X = LEFT_W
RIGHT_W = W - LEFT_W

# ── DATA (from Excel + API, validated) ────────────────────────────────────────
# All metrics sourced from Excel sheets (already computed incrementally)

# CGPO Spring Sale  (qkARQNEM)  — April 2026
SPRING_ROWS = [
    # (model_label, config, cvr_lift, test_aov, aov_delta, iroAS, incr_rev, spend)
    ("Model 338", "4E·20K·10S",  0.6311, 282.50, 0.2147, 21.52, 129179.98, 6003.542),
    ("Model 184", "6E·20K·5S",   0.4917, 237.65, 0.0966, 16.60,  99881.16, 6017.530),
    ("Model 182", "4E·20K·5S",   0.4122, 225.80, 0.0350, 11.52, 138716.09, 12042.592),
    ("Model 276", "2E·20K·10S",  0.1795, 257.07, 0.1965,  7.60,  68524.32, 9010.962),
    ("Model 394", "4E·10K·2S",   0.0923, 226.49, 0.0348,  2.71,  24449.71, 9020.108),
    ("Model 392", "4E·10K·5S",   0.4556, 210.04,-0.2452,  3.51,   9457.91, 2696.456),
    ("Model 72",  "4E·20K·2S",  -0.0638, 241.81, 0.0577, -0.20,  -1784.59, 9008.810),
]
SPRING_TOTALS = dict(
    portfolio_iroAS=8.71,
    total_incr_rev=468424.59,
    total_spend=53800,
    weighted_cvr_lift=0.2626,
    test_reach=100000,
    date="April 2026",
    n_audiences=7,
)

# CG TIB May  (akxzEnAV)  — May 2026
TIB_ROWS = [
    # (model_label, config, cvr_lift, test_aov, aov_delta, iroAS, incr_rev, spend)
    ("Model 399", "2E·10K·10S",   0.9497, 442.27,  0.0724, 44.82, 266921.73, 5955.660),
    ("Model 392", "4E·10K·5S",    0.0527, 370.58,  0.0892,  4.61,  26254.48, 5696.882),
    ("Model 182", "4E·20K·5S",    0.1877, 341.75, -0.1529,  0.24,   1426.71, 5989.016),
    ("Active CRM","Actives L12",  0.0082, 414.18,  0.0107,  1.82,  10473.59, 5756.062),
    ("Lapsed",    "1yr+ Lapsed",  0.0178, 455.28,  0.0333,  0.51,   5236.17, 10227.380),
    ("Model 276", "2E·20K·10S",   None,   268.96, -0.2636,-14.63, -98381.61, 6725.000),
    ("Model 72",  "4E·20K·2S",    None,   564.44,  0.2843, -7.42, -49870.40, 6725.000),
    ("Model 394", "4E·10K·2S",    None,   444.58,  0.3072, -8.53, -57346.58, 6725.000),
]
TIB_TOTALS = dict(
    portfolio_iroAS=1.95,
    total_incr_rev=104714.09,
    total_spend=53800,
    weighted_cvr_lift=None,   # mixed (some lift=0/'-')
    test_reach=100000,
    date="May 2026",
    n_audiences=8,
)

# CGPO Summer Sale  (ykjlvBOV)  — Summer 2026
SUMMER_ROWS = [
    # (model_label, config, cvr_lift, test_aov, aov_delta, iroAS, incr_rev, spend)
    ("Model 338", "4E·20K·10S",  2.1453, 219.87, -0.0000, 17.96, 152396.30, 8485.500),
    ("Model 184", "6E·20K·5S",   0.9760, 203.30,  0.0074, 11.04, 124945.59, 11314.000),
    ("Model 399", "2E·10K·10S",  0.2599, 245.07,  0.1723,  5.59,  63202.05, 11314.000),
    ("Model 400", "4E·10K·10S",  0.0583, 312.98,  0.3696,  3.11,  26390.74, 8485.500),
    ("Model 529", "8E·10K·5S",   0.1824, 359.94,  0.3269,  2.11,  17884.16, 8485.500),
    ("Model 392", "4E·10K·5S",   0.3456, 210.62, -0.2128,  1.06,   8988.08, 8485.500),
]
SUMMER_TOTALS = dict(
    portfolio_iroAS=6.96,
    total_incr_rev=393806.93,
    total_spend=56570,
    weighted_cvr_lift=0.7304,
    test_reach=100000,
    date="Summer 2026",
    n_audiences=6,
)

# ── Helper Functions ──────────────────────────────────────────────────────────

def rgb_hex(r):
    """Return #RRGGBB string from RGBColor."""
    return "#{:02X}{:02X}{:02X}".format(r[0], r[1], r[2])

def add_rect(slide, left, top, width, height, fill_rgb=None, line_rgb=None):
    """Add a filled rectangle shape."""
    shape = slide.shapes.add_shape(
        1,  # MSO_SHAPE_TYPE.RECTANGLE
        left, top, width, height
    )
    shape.line.fill.background()
    if fill_rgb:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_rgb
    else:
        shape.fill.background()
    if line_rgb:
        shape.line.color.rgb = line_rgb
    else:
        shape.line.fill.background()
    return shape

def add_text_box(slide, left, top, width, height, text, font_size=12,
                 bold=False, color=WHITE, align=PP_ALIGN.LEFT, wrap=True,
                 italic=False, line_spacing=None):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    txBox.word_wrap = wrap
    tf = txBox.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.italic = italic
    run.font.name = "Arial"
    if line_spacing:
        from pptx.util import Pt as _Pt
        p.line_spacing = _Pt(line_spacing)
    return txBox

def add_stat_block(slide, left, top, label, value, value_size=36, label_size=11):
    """Add a big teal number + label below it."""
    # Value
    add_text_box(slide, left, top, Inches(3.8), Inches(0.55),
                 value, font_size=value_size, bold=True, color=TEAL,
                 align=PP_ALIGN.LEFT)
    # Label
    add_text_box(slide, left, top + Inches(0.52), Inches(3.8), Inches(0.28),
                 label, font_size=label_size, bold=False, color=WHITE,
                 align=PP_ALIGN.LEFT)

def fmt_iroAS(v):
    if v is None: return "N/A"
    return f"{v:.1f}x"

def fmt_rev(v):
    if v is None: return "N/A"
    if abs(v) >= 1e6:
        return f"${v/1e6:.2f}M"
    elif abs(v) >= 1e3:
        return f"${v/1e3:.0f}K"
    else:
        return f"${v:.0f}"

def fmt_pct(v, plus=True):
    if v is None: return "N/A"
    s = f"{v*100:+.1f}%" if plus else f"{v*100:.1f}%"
    return s

def fmt_aov(v):
    if v is None: return "N/A"
    return f"${v:.0f}"

def row_color(iroAS_val):
    if iroAS_val is None:
        return AMBER_TINT, DARK_TEXT
    if iroAS_val >= 10:
        return TEAL, WHITE
    elif iroAS_val >= 5:
        return TEAL_LIGHT, DARK_TEXT
    elif iroAS_val < 2:
        return AMBER_TINT, DARK_TEXT
    else:
        return WHITE, DARK_TEXT


def build_campaign_slide(rows, totals, campaign_name, brand_name, date_range, file_path):
    """Build a single-slide campaign one-sheeter PPTX."""
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    blank_layout = prs.slide_layouts[6]  # completely blank
    slide = prs.slides.add_slide(blank_layout)

    # ── Left Panel (navy background) ──────────────────────────────────────────
    add_rect(slide, 0, 0, LEFT_W, H, fill_rgb=NAVY)

    # Brand logo text
    add_text_box(slide, Inches(0.25), Inches(0.28), LEFT_W - Inches(0.3), Inches(0.32),
                 brand_name.upper(), font_size=9, bold=True, color=TEAL,
                 align=PP_ALIGN.LEFT)

    # Campaign name
    add_text_box(slide, Inches(0.25), Inches(0.62), LEFT_W - Inches(0.3), Inches(1.1),
                 campaign_name, font_size=22, bold=True, color=WHITE,
                 align=PP_ALIGN.LEFT)

    # Subtitle
    add_text_box(slide, Inches(0.25), Inches(1.75), LEFT_W - Inches(0.3), Inches(0.30),
                 "Postie CRM Optimization", font_size=10, bold=False, color=TEAL,
                 align=PP_ALIGN.LEFT)

    # Divider line
    line = slide.shapes.add_shape(1, Inches(0.25), Inches(2.18), LEFT_W - Inches(0.5), Pt(1))
    line.fill.solid(); line.fill.fore_color.rgb = TEAL
    line.line.fill.background()

    # Stat blocks
    stat_top = Inches(2.40)
    stat_gap = Inches(1.35)

    add_stat_block(slide, Inches(0.25), stat_top,
                   "Portfolio iROAS",
                   f"{totals['portfolio_iroAS']:.1f}x iROAS",
                   value_size=30)

    add_stat_block(slide, Inches(0.25), stat_top + stat_gap,
                   "Total Incremental Revenue",
                   fmt_rev(totals['total_incr_rev']),
                   value_size=30)

    # CVR lift
    if totals.get('weighted_cvr_lift') is not None:
        lift_str = f"+{totals['weighted_cvr_lift']*100:.0f}% CVR Lift"
    else:
        lift_str = "Mixed CVR Lift"

    add_stat_block(slide, Inches(0.25), stat_top + stat_gap * 2,
                   "Portfolio CVR Lift",
                   lift_str,
                   value_size=30)

    # Date + attribution note at bottom of left panel
    add_text_box(slide, Inches(0.25), H - Inches(0.70), LEFT_W - Inches(0.3), Inches(0.35),
                 date_range, font_size=9, bold=False, color=STEEL,
                 align=PP_ALIGN.LEFT)
    add_text_box(slide, Inches(0.25), H - Inches(0.42), LEFT_W - Inches(0.3), Inches(0.35),
                 f"Spend: ${totals['total_spend']:,.0f}  |  Reach: {totals['test_reach']:,}",
                 font_size=8, bold=False, color=STEEL, align=PP_ALIGN.LEFT)

    # ── Right Panel (white background) ────────────────────────────────────────
    add_rect(slide, RIGHT_X, 0, RIGHT_W, H, fill_rgb=WHITE)

    # Header
    add_text_box(slide, RIGHT_X + Inches(0.25), Inches(0.22), RIGHT_W - Inches(0.4), Inches(0.35),
                 "Audience Performance", font_size=14, bold=True, color=DARK_TEXT,
                 align=PP_ALIGN.LEFT)

    # Table layout
    col_labels = ["Model", "Config", "CVR Lift", "Test AOV", "AOV Δ", "iROAS", "Incr Rev"]
    col_widths  = [Inches(1.10), Inches(1.20), Inches(0.95), Inches(1.00), Inches(0.90), Inches(0.90), Inches(1.05)]
    # Ensure columns fit; total right panel width is ~8.65"
    table_left = RIGHT_X + Inches(0.22)
    table_top  = Inches(0.65)
    row_h      = Inches(0.485)
    header_h   = Inches(0.42)

    # Header row
    x = table_left
    for i, (lbl, cw) in enumerate(zip(col_labels, col_widths)):
        add_rect(slide, x, table_top, cw, header_h, fill_rgb=NAVY)
        add_text_box(slide, x + Pt(4), table_top + Pt(5), cw - Pt(6), header_h - Pt(4),
                     lbl, font_size=9, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        x += cw

    # Sort rows by iROAS descending (treat None as -9999)
    sorted_rows = sorted(rows, key=lambda r: r[5] if r[5] is not None else -9999, reverse=True)

    for ri, row in enumerate(sorted_rows):
        model_lbl, config, cvr_lift, test_aov, aov_delta, iroAS, incr_rev, spend = row
        row_top = table_top + header_h + ri * row_h

        bg_color, txt_color = row_color(iroAS)

        x = table_left
        cells = [
            model_lbl,
            config,
            fmt_pct(cvr_lift) if cvr_lift is not None else "—",
            fmt_aov(test_aov),
            fmt_pct(aov_delta) if aov_delta is not None else "—",
            fmt_iroAS(iroAS),
            fmt_rev(incr_rev),
        ]
        for ci, (cell_val, cw) in enumerate(zip(cells, col_widths)):
            add_rect(slide, x, row_top, cw, row_h, fill_rgb=bg_color)
            c_bold = (ci == 5)  # bold iROAS column
            add_text_box(slide, x + Pt(4), row_top + Pt(4), cw - Pt(6), row_h - Pt(4),
                         cell_val, font_size=9.5, bold=c_bold, color=txt_color,
                         align=PP_ALIGN.CENTER)
            x += cw

    # Portfolio total row
    n_data_rows = len(sorted_rows)
    total_top = table_top + header_h + n_data_rows * row_h
    add_rect(slide, table_left, total_top, sum(col_widths), row_h, fill_rgb=DARK_TEXT)
    total_cells = [
        "TOTAL", "",
        fmt_pct(totals.get('weighted_cvr_lift')) if totals.get('weighted_cvr_lift') is not None else "—",
        "", "",
        fmt_iroAS(totals['portfolio_iroAS']),
        fmt_rev(totals['total_incr_rev']),
    ]
    x = table_left
    for ci, (cell_val, cw) in enumerate(zip(total_cells, col_widths)):
        add_text_box(slide, x + Pt(4), total_top + Pt(4), cw - Pt(6), row_h - Pt(4),
                     cell_val, font_size=9.5, bold=True, color=WHITE,
                     align=PP_ALIGN.CENTER)
        x += cw

    # Footer note
    footer_top = H - Inches(0.38)
    add_rect(slide, RIGHT_X, footer_top - Inches(0.04), RIGHT_W, Inches(0.42), fill_rgb=ICE_BLUE)
    add_text_box(slide, RIGHT_X + Inches(0.22), footer_top, RIGHT_W - Inches(0.4), Inches(0.35),
                 "Volume-weighted portfolio metrics.  Incr Rev = Test Rev − (Test Reach / Ctrl Reach) × Ctrl Rev  |  Full 29-day attribution window",
                 font_size=7.5, bold=False, color=STEEL, align=PP_ALIGN.LEFT)

    prs.save(file_path)
    print(f"  Saved: {file_path}")


def build_brand_summary_slide(brand_name, subtitle, campaigns, file_path):
    """
    campaigns: list of dicts with keys:
      name, date, portfolio_iroAS, total_incr_rev, total_spend, n_audiences
    """
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # ── Left Panel ────────────────────────────────────────────────────────────
    add_rect(slide, 0, 0, LEFT_W, H, fill_rgb=NAVY)

    add_text_box(slide, Inches(0.25), Inches(0.28), LEFT_W - Inches(0.3), Inches(0.32),
                 "POSTIE CRM OPTIMIZATION", font_size=9, bold=True, color=TEAL,
                 align=PP_ALIGN.LEFT)

    add_text_box(slide, Inches(0.25), Inches(0.62), LEFT_W - Inches(0.3), Inches(0.72),
                 brand_name, font_size=22, bold=True, color=WHITE,
                 align=PP_ALIGN.LEFT)

    add_text_box(slide, Inches(0.25), Inches(1.40), LEFT_W - Inches(0.3), Inches(0.30),
                 subtitle, font_size=11, bold=False, color=TEAL,
                 align=PP_ALIGN.LEFT)

    line = slide.shapes.add_shape(1, Inches(0.25), Inches(1.82), LEFT_W - Inches(0.5), Pt(1))
    line.fill.solid(); line.fill.fore_color.rgb = TEAL
    line.line.fill.background()

    # Cumulative metrics
    total_incr_rev = sum(c['total_incr_rev'] for c in campaigns)
    total_spend = sum(c['total_spend'] for c in campaigns)
    cum_iroAS = total_incr_rev / total_spend if total_spend else 0

    stat_top = Inches(2.00)
    stat_gap = Inches(1.35)
    add_stat_block(slide, Inches(0.25), stat_top,
                   "Cumulative iROAS",
                   f"{cum_iroAS:.1f}x iROAS",
                   value_size=30)
    add_stat_block(slide, Inches(0.25), stat_top + stat_gap,
                   "Total Incremental Revenue",
                   fmt_rev(total_incr_rev),
                   value_size=30)
    add_stat_block(slide, Inches(0.25), stat_top + stat_gap * 2,
                   "Total Campaigns",
                   f"{len(campaigns)} Campaigns",
                   value_size=30)

    # Date span
    dates = [c['date'] for c in campaigns]
    add_text_box(slide, Inches(0.25), H - Inches(0.55), LEFT_W - Inches(0.3), Inches(0.35),
                 " | ".join(dates), font_size=9, bold=False, color=STEEL,
                 align=PP_ALIGN.LEFT)

    # ── Right Panel ───────────────────────────────────────────────────────────
    add_rect(slide, RIGHT_X, 0, RIGHT_W, H, fill_rgb=WHITE)

    add_text_box(slide, RIGHT_X + Inches(0.25), Inches(0.22), RIGHT_W - Inches(0.4), Inches(0.35),
                 "Campaign Performance Summary", font_size=14, bold=True, color=DARK_TEXT,
                 align=PP_ALIGN.LEFT)

    # Summary table
    col_labels = ["Campaign", "Date", "Portfolio iROAS", "Total Incr Rev", "Total Spend", "# Audiences"]
    col_widths  = [Inches(2.50), Inches(1.10), Inches(1.50), Inches(1.50), Inches(1.30), Inches(1.10)]
    table_left = RIGHT_X + Inches(0.22)
    table_top  = Inches(0.65)
    row_h      = Inches(0.60)
    header_h   = Inches(0.42)

    x = table_left
    for lbl, cw in zip(col_labels, col_widths):
        add_rect(slide, x, table_top, cw, header_h, fill_rgb=NAVY)
        add_text_box(slide, x + Pt(4), table_top + Pt(5), cw - Pt(6), header_h - Pt(4),
                     lbl, font_size=9, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        x += cw

    for ri, camp in enumerate(campaigns):
        row_top = table_top + header_h + ri * row_h
        bg = ICE_BLUE if ri % 2 == 1 else WHITE
        x = table_left
        cells = [
            camp['name'],
            camp['date'],
            fmt_iroAS(camp['portfolio_iroAS']),
            fmt_rev(camp['total_incr_rev']),
            f"${camp['total_spend']:,.0f}",
            str(camp['n_audiences']),
        ]
        for ci, (cell_val, cw) in enumerate(zip(cells, col_widths)):
            add_rect(slide, x, row_top, cw, row_h, fill_rgb=bg)
            c_bold = (ci == 2 and camp['portfolio_iroAS'] >= 5)
            add_text_box(slide, x + Pt(4), row_top + Pt(6), cw - Pt(6), row_h - Pt(6),
                         cell_val, font_size=9.5, bold=c_bold, color=DARK_TEXT,
                         align=PP_ALIGN.CENTER)
            x += cw

    # TOTAL row
    total_top = table_top + header_h + len(campaigns) * row_h
    add_rect(slide, table_left, total_top, sum(col_widths), row_h, fill_rgb=DARK_TEXT)
    total_cells = [
        "TOTAL", "",
        fmt_iroAS(cum_iroAS),
        fmt_rev(total_incr_rev),
        f"${total_spend:,.0f}",
        str(sum(c['n_audiences'] for c in campaigns)),
    ]
    x = table_left
    for cell_val, cw in zip(total_cells, col_widths):
        add_text_box(slide, x + Pt(4), total_top + Pt(6), cw - Pt(6), row_h - Pt(6),
                     cell_val, font_size=9.5, bold=True, color=WHITE,
                     align=PP_ALIGN.CENTER)
        x += cw

    # Compounding revenue callout box
    callout_top = total_top + row_h + Inches(0.25)
    callout_h = H - callout_top - Inches(0.18)
    add_rect(slide, RIGHT_X + Inches(0.22), callout_top,
             RIGHT_W - Inches(0.44), callout_h,
             fill_rgb=ICE_BLUE)

    # Build compounding string
    parts = [f"{fmt_rev(c['total_incr_rev'])} ({c['name']})" for c in campaigns]
    compound_str = " + ".join(parts) + f" = {fmt_rev(total_incr_rev)} total {brand_name} incremental revenue"

    add_text_box(slide, RIGHT_X + Inches(0.38), callout_top + Inches(0.12),
                 RIGHT_W - Inches(0.70), callout_h - Inches(0.18),
                 "Compounding Incremental Revenue",
                 font_size=10, bold=True, color=DARK_TEXT, align=PP_ALIGN.LEFT)
    add_text_box(slide, RIGHT_X + Inches(0.38), callout_top + Inches(0.40),
                 RIGHT_W - Inches(0.70), callout_h - Inches(0.45),
                 compound_str,
                 font_size=9, bold=False, color=DARK_TEXT, align=PP_ALIGN.LEFT, wrap=True)

    # Footer
    footer_top = H - Inches(0.38)
    add_rect(slide, RIGHT_X, footer_top - Inches(0.04), RIGHT_W, Inches(0.42), fill_rgb=ICE_BLUE)
    add_text_box(slide, RIGHT_X + Inches(0.22), footer_top, RIGHT_W - Inches(0.4), Inches(0.35),
                 "Volume-weighted portfolio metrics.  Incr Rev = Test Rev − (Test Reach / Ctrl Reach) × Ctrl Rev  |  Full 29-day attribution window",
                 font_size=7.5, bold=False, color=STEEL, align=PP_ALIGN.LEFT)

    prs.save(file_path)
    print(f"  Saved: {file_path}")


# ── BUILD PPTX ────────────────────────────────────────────────────────────────
print("Building PPTX one-sheeters...")

# 1. CGPO Spring Sale
build_campaign_slide(
    rows=SPRING_ROWS,
    totals=SPRING_TOTALS,
    campaign_name="CGPO Spring Sale 2026",
    brand_name="Callaway Golf Pre-Owned",
    date_range="April 2026  |  Campaign: qkARQNEM",
    file_path=os.path.join(PPTX_DIR, "CGPO_Spring_Sale_2026.pptx"),
)

# 2. CG TIB May
build_campaign_slide(
    rows=TIB_ROWS,
    totals=TIB_TOTALS,
    campaign_name="CG Trade-In Bonus May 2026",
    brand_name="Callaway Golf",
    date_range="May 2026  |  Campaign: akxzEnAV",
    file_path=os.path.join(PPTX_DIR, "CG_TIB_May_2026.pptx"),
)

# 3. CGPO Summer Sale
build_campaign_slide(
    rows=SUMMER_ROWS,
    totals=SUMMER_TOTALS,
    campaign_name="CGPO Summer Sale 2026",
    brand_name="Callaway Golf Pre-Owned",
    date_range="Summer 2026  |  Campaign: ykjlvBOV",
    file_path=os.path.join(PPTX_DIR, "CGPO_Summer_Sale_2026.pptx"),
)

# 4. CG Brand Summary
build_brand_summary_slide(
    brand_name="Callaway Golf",
    subtitle="Program Summary — Spring/Summer 2026",
    campaigns=[
        dict(name="CG TIB May 2026", date="May 2026",
             portfolio_iroAS=TIB_TOTALS['portfolio_iroAS'],
             total_incr_rev=TIB_TOTALS['total_incr_rev'],
             total_spend=TIB_TOTALS['total_spend'],
             n_audiences=TIB_TOTALS['n_audiences']),
    ],
    file_path=os.path.join(PPTX_DIR, "CG_Program_Summary.pptx"),
)

# 5. CGPO Brand Summary
build_brand_summary_slide(
    brand_name="Callaway Golf Pre-Owned",
    subtitle="Program Summary — Spring/Summer 2026",
    campaigns=[
        dict(name="CGPO Spring Sale 2026", date="April 2026",
             portfolio_iroAS=SPRING_TOTALS['portfolio_iroAS'],
             total_incr_rev=SPRING_TOTALS['total_incr_rev'],
             total_spend=SPRING_TOTALS['total_spend'],
             n_audiences=SPRING_TOTALS['n_audiences']),
        dict(name="CGPO Summer Sale 2026", date="Summer 2026",
             portfolio_iroAS=SUMMER_TOTALS['portfolio_iroAS'],
             total_incr_rev=SUMMER_TOTALS['total_incr_rev'],
             total_spend=SUMMER_TOTALS['total_spend'],
             n_audiences=SUMMER_TOTALS['n_audiences']),
    ],
    file_path=os.path.join(PPTX_DIR, "CGPO_Program_Summary.pptx"),
)

print("\nAll PPTX files built.")

# ── COMPUTE CGPO cumulative for HTML ─────────────────────────────────────────
cgpo_total_incr = SPRING_TOTALS['total_incr_rev'] + SUMMER_TOTALS['total_incr_rev']
cgpo_total_spend = SPRING_TOTALS['total_spend'] + SUMMER_TOTALS['total_spend']
cgpo_cum_iroAS = cgpo_total_incr / cgpo_total_spend

cg_total_incr = TIB_TOTALS['total_incr_rev']
cg_total_spend = TIB_TOTALS['total_spend']
cg_cum_iroAS = cg_total_incr / cg_total_spend

# ── BUILD HTML DASHBOARD ─────────────────────────────────────────────────────

html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Callaway Golf × Postie | Performance Dashboard</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<style>
  @import url('https://fonts.googleapis.com/css2?family=Lato:wght@300;400;700;900&display=swap');

  :root {{
    --navy:       #163D5D;
    --teal:       #00C5C1;
    --ice-blue:   #EAF8FF;
    --light-blue: #CADCFC;
    --steel:      #5A7A8A;
    --amber:      #F0B840;
    --white:      #FFFFFF;
    --dark-text:  #1E4768;
    --card-shadow: 0 2px 12px rgba(22,61,93,0.10);
    --card-shadow-hover: 0 8px 28px rgba(22,61,93,0.18);
  }}

  *, *::before, *::after {{ box-sizing: border-box; margin: 0; padding: 0; }}

  body {{
    font-family: 'Lato', Arial, sans-serif;
    background: var(--ice-blue);
    color: var(--dark-text);
    min-height: 100vh;
  }}

  /* ── Header ── */
  .site-header {{
    background: var(--navy);
    padding: 0 40px;
    display: flex;
    align-items: center;
    justify-content: space-between;
    height: 72px;
    box-shadow: 0 2px 8px rgba(0,0,0,0.25);
  }}
  .site-header h1 {{
    color: var(--white);
    font-size: 1.35rem;
    font-weight: 700;
    letter-spacing: 0.01em;
  }}
  .site-header h1 span {{ color: var(--teal); }}
  .header-badge {{
    background: var(--teal);
    color: var(--white);
    font-size: 0.72rem;
    font-weight: 700;
    padding: 4px 12px;
    border-radius: 20px;
    letter-spacing: 0.06em;
    text-transform: uppercase;
  }}

  /* ── Main container ── */
  .container {{
    max-width: 1280px;
    margin: 0 auto;
    padding: 40px 32px 60px;
  }}

  /* ── Brand Section ── */
  .brand-section {{
    margin-bottom: 56px;
  }}
  .brand-header {{
    display: flex;
    align-items: center;
    gap: 16px;
    margin-bottom: 24px;
  }}
  .brand-pill {{
    background: var(--navy);
    color: var(--white);
    font-size: 0.80rem;
    font-weight: 700;
    padding: 4px 14px;
    border-radius: 20px;
    text-transform: uppercase;
    letter-spacing: 0.08em;
  }}
  .brand-title {{
    font-size: 1.55rem;
    font-weight: 900;
    color: var(--navy);
  }}
  .brand-divider {{
    width: 100%;
    height: 2px;
    background: linear-gradient(to right, var(--teal), transparent);
    margin-bottom: 24px;
  }}

  /* ── Campaign grid ── */
  .campaign-grid {{
    display: grid;
    grid-template-columns: repeat(auto-fill, minmax(340px, 1fr));
    gap: 20px;
    margin-bottom: 24px;
  }}

  /* ── Campaign Card ── */
  .campaign-card {{
    background: var(--white);
    border-radius: 14px;
    box-shadow: var(--card-shadow);
    padding: 24px 26px 20px;
    display: flex;
    flex-direction: column;
    gap: 0;
    transition: box-shadow 0.2s ease, transform 0.2s ease;
    border-top: 4px solid var(--teal);
  }}
  .campaign-card:hover {{
    box-shadow: var(--card-shadow-hover);
    transform: translateY(-3px);
  }}
  .card-top-row {{
    display: flex;
    justify-content: space-between;
    align-items: flex-start;
    margin-bottom: 4px;
  }}
  .card-title {{
    font-size: 1.05rem;
    font-weight: 700;
    color: var(--navy);
    line-height: 1.3;
    flex: 1;
    padding-right: 12px;
  }}
  .card-id {{
    font-size: 0.68rem;
    color: var(--steel);
    font-weight: 400;
    background: var(--ice-blue);
    padding: 2px 8px;
    border-radius: 8px;
    white-space: nowrap;
    margin-top: 3px;
  }}
  .card-date {{
    font-size: 0.78rem;
    color: var(--steel);
    margin-bottom: 16px;
    font-weight: 400;
  }}

  /* ── Metric Pills ── */
  .metrics-row {{
    display: flex;
    gap: 8px;
    flex-wrap: wrap;
    margin-bottom: 18px;
  }}
  .metric-pill {{
    background: var(--ice-blue);
    border: 1px solid var(--light-blue);
    border-radius: 999px;
    padding: 5px 13px;
    display: flex;
    flex-direction: column;
    align-items: center;
    min-width: 90px;
    flex: 1;
  }}
  .metric-pill .pill-value {{
    font-size: 1.08rem;
    font-weight: 900;
    color: var(--teal);
    line-height: 1.2;
    white-space: nowrap;
  }}
  .metric-pill .pill-label {{
    font-size: 0.63rem;
    font-weight: 600;
    color: var(--steel);
    text-transform: uppercase;
    letter-spacing: 0.05em;
    margin-top: 1px;
  }}
  .metric-pill.highlight {{
    background: var(--navy);
    border-color: var(--navy);
  }}
  .metric-pill.highlight .pill-value {{ color: var(--teal); }}
  .metric-pill.highlight .pill-label {{ color: rgba(255,255,255,0.60); }}

  /* ── Download button ── */
  .btn-download {{
    display: inline-flex;
    align-items: center;
    gap: 7px;
    background: var(--teal);
    color: var(--white);
    text-decoration: none;
    font-size: 0.82rem;
    font-weight: 700;
    padding: 9px 18px;
    border-radius: 8px;
    transition: background 0.18s, transform 0.15s;
    align-self: flex-start;
    margin-top: auto;
    letter-spacing: 0.02em;
  }}
  .btn-download:hover {{
    background: var(--navy);
    transform: translateY(-1px);
  }}
  .btn-download svg {{
    flex-shrink: 0;
  }}

  /* ── Brand Aggregate Card ── */
  .brand-agg-card {{
    background: var(--navy);
    border-radius: 16px;
    padding: 30px 34px;
    display: flex;
    align-items: center;
    justify-content: space-between;
    gap: 32px;
    box-shadow: var(--card-shadow-hover);
    transition: transform 0.2s ease;
    flex-wrap: wrap;
  }}
  .brand-agg-card:hover {{
    transform: translateY(-3px);
  }}
  .agg-left {{ flex: 1; min-width: 260px; }}
  .agg-tag {{
    font-size: 0.72rem;
    font-weight: 700;
    color: var(--teal);
    text-transform: uppercase;
    letter-spacing: 0.10em;
    margin-bottom: 6px;
  }}
  .agg-brand-name {{
    font-size: 1.30rem;
    font-weight: 900;
    color: var(--white);
    margin-bottom: 4px;
  }}
  .agg-subtitle {{
    font-size: 0.80rem;
    color: rgba(255,255,255,0.55);
    margin-bottom: 18px;
  }}
  .agg-campaign-list {{
    list-style: none;
    display: flex;
    flex-direction: column;
    gap: 3px;
  }}
  .agg-campaign-list li {{
    font-size: 0.78rem;
    color: rgba(255,255,255,0.75);
    display: flex;
    align-items: center;
    gap: 7px;
  }}
  .agg-campaign-list li::before {{
    content: '';
    display: inline-block;
    width: 6px;
    height: 6px;
    border-radius: 50%;
    background: var(--teal);
    flex-shrink: 0;
  }}
  .agg-metrics {{
    display: flex;
    gap: 32px;
    align-items: center;
    flex-wrap: wrap;
  }}
  .agg-metric {{
    display: flex;
    flex-direction: column;
    align-items: center;
  }}
  .agg-metric-value {{
    font-size: 2.20rem;
    font-weight: 900;
    color: var(--teal);
    line-height: 1.1;
    white-space: nowrap;
  }}
  .agg-metric-label {{
    font-size: 0.72rem;
    font-weight: 600;
    color: rgba(255,255,255,0.55);
    text-transform: uppercase;
    letter-spacing: 0.08em;
    margin-top: 3px;
  }}
  .agg-divider {{
    width: 1px;
    height: 56px;
    background: rgba(255,255,255,0.18);
  }}
  .agg-right {{
    display: flex;
    align-items: center;
  }}
  .btn-download-white {{
    display: inline-flex;
    align-items: center;
    gap: 7px;
    background: rgba(255,255,255,0.10);
    border: 2px solid var(--teal);
    color: var(--white);
    text-decoration: none;
    font-size: 0.84rem;
    font-weight: 700;
    padding: 10px 20px;
    border-radius: 10px;
    transition: background 0.18s, transform 0.15s;
    white-space: nowrap;
    letter-spacing: 0.02em;
  }}
  .btn-download-white:hover {{
    background: var(--teal);
    transform: translateY(-1px);
  }}

  /* ── Footer ── */
  footer {{
    background: var(--navy);
    text-align: center;
    padding: 22px 0;
    color: rgba(255,255,255,0.45);
    font-size: 0.78rem;
    letter-spacing: 0.04em;
  }}
  footer a {{
    color: var(--teal);
    text-decoration: none;
  }}
  footer a:hover {{ text-decoration: underline; }}

  @media (max-width: 720px) {{
    .container {{ padding: 24px 16px 40px; }}
    .site-header {{ padding: 0 20px; }}
    .site-header h1 {{ font-size: 1rem; }}
    .agg-metrics {{ gap: 18px; }}
    .brand-agg-card {{ padding: 22px 20px; }}
  }}
</style>
</head>
<body>

<!-- ── Site Header ─────────────────────────────────────────────────────── -->
<header class="site-header">
  <h1>Callaway Golf <span>×</span> Postie &nbsp;|&nbsp; Performance Dashboard</h1>
  <span class="header-badge">CRM Optimization</span>
</header>

<!-- ── Main Content ───────────────────────────────────────────────────── -->
<main class="container">

  <!-- ══════════════════════════════════════════════════════════════════ -->
  <!--  CALLAWAY GOLF  (CG)                                             -->
  <!-- ══════════════════════════════════════════════════════════════════ -->
  <section class="brand-section">
    <div class="brand-header">
      <span class="brand-pill">Brand</span>
      <h2 class="brand-title">Callaway Golf</h2>
    </div>
    <div class="brand-divider"></div>

    <!-- Campaign Cards -->
    <div class="campaign-grid">

      <!-- CG TIB May 2026 -->
      <div class="campaign-card">
        <div class="card-top-row">
          <div class="card-title">CG Trade-In Bonus May 2026</div>
          <span class="card-id">akxzEnAV</span>
        </div>
        <div class="card-date">May 2026 &nbsp;·&nbsp; 8 audiences &nbsp;·&nbsp; $53,800 spend</div>
        <div class="metrics-row">
          <div class="metric-pill highlight">
            <span class="pill-value">1.9x</span>
            <span class="pill-label">iROAS</span>
          </div>
          <div class="metric-pill">
            <span class="pill-value">$105K</span>
            <span class="pill-label">Incr Rev</span>
          </div>
          <div class="metric-pill">
            <span class="pill-value">Mixed</span>
            <span class="pill-label">CVR Lift</span>
          </div>
        </div>
        <a class="btn-download" href="one-sheeters/CG_TIB_May_2026.pptx">
          <svg width="14" height="14" viewBox="0 0 14 14" fill="none"><path d="M7 1v8M3.5 6.5 7 10l3.5-3.5M2 12h10" stroke="currentColor" stroke-width="1.6" stroke-linecap="round" stroke-linejoin="round"/></svg>
          Download One-Sheeter
        </a>
      </div>

    </div><!-- /campaign-grid -->

    <!-- CG Brand Aggregate Card -->
    <div class="brand-agg-card">
      <div class="agg-left">
        <div class="agg-tag">Program Summary</div>
        <div class="agg-brand-name">Callaway Golf</div>
        <div class="agg-subtitle">Spring / Summer 2026 &nbsp;·&nbsp; Postie CRM Optimization</div>
        <ul class="agg-campaign-list">
          <li>CG Trade-In Bonus May 2026 (akxzEnAV)</li>
        </ul>
      </div>
      <div class="agg-metrics">
        <div class="agg-metric">
          <span class="agg-metric-value">{cg_cum_iroAS:.1f}x</span>
          <span class="agg-metric-label">Cumulative iROAS</span>
        </div>
        <div class="agg-divider"></div>
        <div class="agg-metric">
          <span class="agg-metric-value">${cg_total_incr/1000:.0f}K</span>
          <span class="agg-metric-label">Total Incr Rev</span>
        </div>
      </div>
      <div class="agg-right">
        <a class="btn-download-white" href="one-sheeters/CG_Program_Summary.pptx">
          <svg width="14" height="14" viewBox="0 0 14 14" fill="none"><path d="M7 1v8M3.5 6.5 7 10l3.5-3.5M2 12h10" stroke="currentColor" stroke-width="1.6" stroke-linecap="round" stroke-linejoin="round"/></svg>
          Download Program Summary
        </a>
      </div>
    </div>
  </section>

  <!-- ══════════════════════════════════════════════════════════════════ -->
  <!--  CALLAWAY GOLF PRE-OWNED  (CGPO)                                 -->
  <!-- ══════════════════════════════════════════════════════════════════ -->
  <section class="brand-section">
    <div class="brand-header">
      <span class="brand-pill">Brand</span>
      <h2 class="brand-title">Callaway Golf Pre-Owned</h2>
    </div>
    <div class="brand-divider"></div>

    <!-- Campaign Cards -->
    <div class="campaign-grid">

      <!-- CGPO Spring Sale 2026 -->
      <div class="campaign-card">
        <div class="card-top-row">
          <div class="card-title">CGPO Spring Sale 2026</div>
          <span class="card-id">qkARQNEM</span>
        </div>
        <div class="card-date">April 2026 &nbsp;·&nbsp; 7 audiences &nbsp;·&nbsp; $53,800 spend</div>
        <div class="metrics-row">
          <div class="metric-pill highlight">
            <span class="pill-value">8.7x</span>
            <span class="pill-label">iROAS</span>
          </div>
          <div class="metric-pill">
            <span class="pill-value">$468K</span>
            <span class="pill-label">Incr Rev</span>
          </div>
          <div class="metric-pill">
            <span class="pill-value">+26.3%</span>
            <span class="pill-label">CVR Lift</span>
          </div>
        </div>
        <a class="btn-download" href="one-sheeters/CGPO_Spring_Sale_2026.pptx">
          <svg width="14" height="14" viewBox="0 0 14 14" fill="none"><path d="M7 1v8M3.5 6.5 7 10l3.5-3.5M2 12h10" stroke="currentColor" stroke-width="1.6" stroke-linecap="round" stroke-linejoin="round"/></svg>
          Download One-Sheeter
        </a>
      </div>

      <!-- CGPO Summer Sale 2026 -->
      <div class="campaign-card">
        <div class="card-top-row">
          <div class="card-title">CGPO Summer Sale 2026</div>
          <span class="card-id">ykjlvBOV</span>
        </div>
        <div class="card-date">Summer 2026 &nbsp;·&nbsp; 6 audiences &nbsp;·&nbsp; $56,570 spend</div>
        <div class="metrics-row">
          <div class="metric-pill highlight">
            <span class="pill-value">7.0x</span>
            <span class="pill-label">iROAS</span>
          </div>
          <div class="metric-pill">
            <span class="pill-value">$394K</span>
            <span class="pill-label">Incr Rev</span>
          </div>
          <div class="metric-pill">
            <span class="pill-value">+73.0%</span>
            <span class="pill-label">CVR Lift</span>
          </div>
        </div>
        <a class="btn-download" href="one-sheeters/CGPO_Summer_Sale_2026.pptx">
          <svg width="14" height="14" viewBox="0 0 14 14" fill="none"><path d="M7 1v8M3.5 6.5 7 10l3.5-3.5M2 12h10" stroke="currentColor" stroke-width="1.6" stroke-linecap="round" stroke-linejoin="round"/></svg>
          Download One-Sheeter
        </a>
      </div>

    </div><!-- /campaign-grid -->

    <!-- CGPO Brand Aggregate Card -->
    <div class="brand-agg-card">
      <div class="agg-left">
        <div class="agg-tag">Program Summary</div>
        <div class="agg-brand-name">Callaway Golf Pre-Owned</div>
        <div class="agg-subtitle">Spring / Summer 2026 &nbsp;·&nbsp; Postie CRM Optimization</div>
        <ul class="agg-campaign-list">
          <li>CGPO Spring Sale 2026 (qkARQNEM) — April 2026</li>
          <li>CGPO Summer Sale 2026 (ykjlvBOV) — Summer 2026</li>
        </ul>
      </div>
      <div class="agg-metrics">
        <div class="agg-metric">
          <span class="agg-metric-value">{cgpo_cum_iroAS:.1f}x</span>
          <span class="agg-metric-label">Cumulative iROAS</span>
        </div>
        <div class="agg-divider"></div>
        <div class="agg-metric">
          <span class="agg-metric-value">${cgpo_total_incr/1000:.0f}K</span>
          <span class="agg-metric-label">Total Incr Rev</span>
        </div>
        <div class="agg-divider"></div>
        <div class="agg-metric">
          <span class="agg-metric-value">$468K + $394K</span>
          <span class="agg-metric-label">Spring + Summer</span>
        </div>
      </div>
      <div class="agg-right">
        <a class="btn-download-white" href="one-sheeters/CGPO_Program_Summary.pptx">
          <svg width="14" height="14" viewBox="0 0 14 14" fill="none"><path d="M7 1v8M3.5 6.5 7 10l3.5-3.5M2 12h10" stroke="currentColor" stroke-width="1.6" stroke-linecap="round" stroke-linejoin="round"/></svg>
          Download Program Summary
        </a>
      </div>
    </div>
  </section>

</main>

<!-- ── Footer ─────────────────────────────────────────────────────────── -->
<footer>
  Powered by <a href="https://postie.com" target="_blank">Postie</a> &nbsp;|&nbsp;
  postie.com &nbsp;|&nbsp; CRM Optimization Program &nbsp;|&nbsp; Dashboard generated October 2026
</footer>

</body>
</html>
"""

html_path = os.path.join(OUTPUT_DIR, "index.html")
with open(html_path, "w", encoding="utf-8") as f:
    f.write(html)
print(f"\nHTML dashboard saved: {html_path}")

# ── Verification Summary ──────────────────────────────────────────────────────
print("\n" + "="*70)
print("KEY METRICS USED IN DASHBOARD")
print("="*70)
print(f"\nCGPO Spring Sale (qkARQNEM) — April 2026:")
print(f"  Portfolio iROAS:   {SPRING_TOTALS['portfolio_iroAS']:.2f}x")
print(f"  Total Incr Rev:    ${SPRING_TOTALS['total_incr_rev']:,.2f}")
print(f"  Weighted CVR Lift: {SPRING_TOTALS['weighted_cvr_lift']*100:.2f}%")
print(f"  Total Spend:       ${SPRING_TOTALS['total_spend']:,.2f}")
print(f"\nCG TIB May (akxzEnAV) — May 2026:")
print(f"  Portfolio iROAS:   {TIB_TOTALS['portfolio_iroAS']:.2f}x")
print(f"  Total Incr Rev:    ${TIB_TOTALS['total_incr_rev']:,.2f}")
print(f"  Weighted CVR Lift: Mixed / not single-valued")
print(f"  Total Spend:       ${TIB_TOTALS['total_spend']:,.2f}")
print(f"\nCGPO Summer Sale (ykjlvBOV) — Summer 2026:")
print(f"  Portfolio iROAS:   {SUMMER_TOTALS['portfolio_iroAS']:.2f}x")
print(f"  Total Incr Rev:    ${SUMMER_TOTALS['total_incr_rev']:,.2f}")
print(f"  Weighted CVR Lift: {SUMMER_TOTALS['weighted_cvr_lift']*100:.2f}%")
print(f"  Total Spend:       ${SUMMER_TOTALS['total_spend']:,.2f}")
print(f"\n── Brand Aggregates ──")
print(f"CG  → iROAS {cg_cum_iroAS:.2f}x | Incr Rev ${cg_total_incr:,.2f}")
print(f"CGPO→ iROAS {cgpo_cum_iroAS:.2f}x | Incr Rev ${cgpo_total_incr:,.2f}")
print(f"\n── Files Created ──")
for fn in sorted(os.listdir(PPTX_DIR)):
    print(f"  {PPTX_DIR}/{fn}")
print(f"  {html_path}")
